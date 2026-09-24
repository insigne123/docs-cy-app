"""Reporte mensual de contratos: estado de avance, vigencia y alertas. Formatos XLSX y PDF."""
from __future__ import annotations

import calendar
import io
from datetime import date, datetime
from typing import Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.contrato import Garantia
from app.services.alertas import calcular_alertas
from app.services.dashboard import MESES_ES, construir_dashboard

MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _fin_de_mes(anio: int, mes: int) -> date:
    return date(anio, mes, calendar.monthrange(anio, mes)[1])


def _clasificar_garantia(venc: date, corte: date, umbral: int) -> tuple[str, int]:
    d = (venc - corte).days
    if d < 0:
        return "vencida", d
    if d <= umbral:
        return "por_vencer", d
    return "vigente", d


def datos_reporte(session: Session, anio: int, mes: int, hoy: Optional[date] = None) -> dict:
    corte = _fin_de_mes(anio, mes)
    periodo = f"{anio:04d}-{mes:02d}"
    etiqueta = f"{MESES_ES[mes].capitalize()} {anio}"

    tablero = construir_dashboard(session, agrupacion="vencimiento", hoy=corte)
    detalle = [c for g in tablero["grupos"] for c in g["contratos"]]
    vencen = next(
        (g["contratos"] for g in tablero["grupos"] if g["periodo"] == periodo), []
    )
    alertas = calcular_alertas(session, hoy=corte)

    cod = {c["id"]: c["codigo"] for c in detalle}
    garantias = []
    for g in session.scalars(select(Garantia).where(Garantia.entidad_tipo == "contrato")):
        clasif, d = _clasificar_garantia(g.fecha_vencimiento, corte, 90)
        garantias.append(
            {
                "contrato": cod.get(g.entidad_id, f"id {g.entidad_id}"),
                "tipo": g.tipo.value,
                "instrumento": g.instrumento.value,
                "emisor": g.emisor or "",
                "monto": f"{g.monto:.2f}",
                "moneda": g.moneda.value,
                "fecha_vencimiento": g.fecha_vencimiento,
                "estado": g.estado.value,
                "dias": d,
                "clasificacion": clasif,
            }
        )

    return {
        "periodo": periodo,
        "etiqueta": etiqueta,
        "corte": corte,
        "generado": datetime.utcnow(),
        "indicadores": tablero["indicadores"],
        "detalle": detalle,
        "vencen_en_el_mes": vencen,
        "alertas": alertas,
        "garantias": garantias,
    }


# --------------------------------------------------------------------- XLSX
def reporte_xlsx(datos: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    ind = datos["indicadores"]

    ws = wb.create_sheet("Resumen")
    filas = [
        ["Reporte mensual de contratos"],
        ["Período", datos["etiqueta"]],
        ["Fecha de corte", datos["corte"].isoformat()],
        ["Generado", datos["generado"].strftime("%Y-%m-%d %H:%M")],
        [],
        ["Total de contratos", ind["total"]],
        ["Línea A", ind["por_linea"].get("A_regular", 0)],
        ["Línea B", ind["por_linea"].get("B_autogestionado", 0)],
        ["Línea C", ind["por_linea"].get("C_licitacion", 0)],
        ["Vigentes", ind["vigentes"]],
        ["Por vencer", ind["por_vencer"]],
        ["Vencidos", ind["vencidos"]],
        ["En renovación", ind["en_renovacion"]],
        [],
        ["Monto total por moneda"],
    ]
    for m, v in ind["monto_total"].items():
        filas.append([m, v])
    filas += [[], ["Alertas"], ["Total", len(datos["alertas"])]]
    for a in datos["alertas"]:
        filas.append([a["nivel"], a["tipo"], a["contrato_codigo"], a["titulo"]])
    for f in filas:
        ws.append(f)

    dcols = ["Código", "Flujo", "Estado", "Contraparte", "Unidad", "Administrador", "Monto",
             "Moneda", "Ingreso", "Firma", "Fin vigencia", "Días", "Semáforo"]
    for nombre, filas_c in (("Detalle", datos["detalle"]), ("Vencen en el mes", datos["vencen_en_el_mes"])):
        w = wb.create_sheet(nombre)
        w.append(dcols)
        for c in filas_c:
            w.append([
                c["codigo"], c["linea_nombre"], c["estado"], c["contraparte"], c["unidad"],
                c["administrador"], c["monto"], c["moneda"],
                c["fecha_ingreso"].isoformat() if c["fecha_ingreso"] else "",
                c["fecha_firma"].isoformat() if c["fecha_firma"] else "",
                c["fecha_fin_vigencia"].isoformat() if c["fecha_fin_vigencia"] else "",
                c["dias_para_vencer"], c["semaforo"],
            ])

    wa = wb.create_sheet("Alertas")
    wa.append(["Nivel", "Tipo", "Contrato", "Título", "Detalle", "Fecha ref.", "Días"])
    for a in datos["alertas"]:
        wa.append([
            a["nivel"], a["tipo"], a["contrato_codigo"], a["titulo"], a["detalle"],
            a["fecha_referencia"].isoformat() if a["fecha_referencia"] else "", a["dias"],
        ])

    wg = wb.create_sheet("Garantias")
    wg.append(["Contrato", "Tipo", "Instrumento", "Emisor", "Monto", "Moneda", "Vencimiento",
               "Estado", "Días", "Clasificación"])
    for g in datos["garantias"]:
        wg.append([
            g["contrato"], g["tipo"], g["instrumento"], g["emisor"], g["monto"], g["moneda"],
            g["fecha_vencimiento"].isoformat(), g["estado"], g["dias"], g["clasificacion"],
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- PDF
def _lat1(v) -> str:
    s = "" if v is None else str(v)
    return s.replace("—", "-").replace("→", "->").encode("latin-1", "replace").decode("latin-1")


def reporte_pdf(datos: dict) -> bytes:
    from fpdf import FPDF

    ind = datos["indicadores"]
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    def linea(texto: str, alto: float = 6, size: int = 9, estilo: str = "") -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", estilo, size)
        pdf.cell(0, alto, _lat1(texto))
        pdf.ln(alto)

    linea("Reporte mensual de contratos", alto=8, size=14, estilo="B")
    linea(
        f"Período: {datos['etiqueta']}   ·   Corte: {datos['corte']}   ·   "
        f"Generado: {datos['generado'].strftime('%Y-%m-%d %H:%M')}"
    )
    pdf.ln(2)
    linea(
        f"Total {ind['total']}  |  A {ind['por_linea'].get('A_regular', 0)}  "
        f"B {ind['por_linea'].get('B_autogestionado', 0)}  C {ind['por_linea'].get('C_licitacion', 0)}  |  "
        f"Vigentes {ind['vigentes']}  ·  Por vencer {ind['por_vencer']}  ·  "
        f"Vencidos {ind['vencidos']}  ·  En renovación {ind['en_renovacion']}",
        size=8,
    )
    montos = "  ".join(f"{m} {v}" for m, v in ind["monto_total"].items()) or "-"
    linea(f"Monto total: {montos}", size=8)
    pdf.ln(2)

    def tabla(titulo, headers, widths, rows):
        linea(titulo, alto=6, size=10, estilo="B")
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 7)
        for h, w in zip(headers, widths):
            pdf.cell(w, 5, _lat1(h), border=1)
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 7)
        if not rows:
            pdf.set_x(pdf.l_margin)
            pdf.cell(sum(widths), 5, _lat1("(sin datos)"), border=1)
            pdf.ln(5)
        for r in rows:
            pdf.set_x(pdf.l_margin)
            for v, w in zip(r, widths):
                pdf.cell(w, 5, _lat1(v)[:38], border=1)
            pdf.ln(5)
        pdf.ln(3)

    tabla(
        "Contratos - estado de avance y vigencia",
        ["Código", "Flujo", "Estado", "Contraparte", "Monto", "Fin vig.", "Días", "Semáforo"],
        [26, 30, 26, 51, 30, 24, 16, 30],
        [
            [c["codigo"], c["linea_nombre"], c["estado"], c["contraparte"],
             f"{c['moneda'] or ''} {c['monto'] or ''}".strip(),
             c["fecha_fin_vigencia"] or ("indefinida" if c["vigencia_indefinida"] else "-"),
             c["dias_para_vencer"] if c["dias_para_vencer"] is not None else "-",
             c["semaforo"]]
            for c in datos["detalle"]
        ],
    )
    tabla(
        "Alertas (vencidos / por vencer / atrasos)",
        ["Nivel", "Tipo", "Contrato", "Título", "Fecha ref.", "Días"],
        [22, 34, 26, 90, 24, 16],
        [[a["nivel"], a["tipo"], a["contrato_codigo"], a["titulo"], a["fecha_referencia"], a["dias"]]
         for a in datos["alertas"]],
    )
    tabla(
        "Garantías",
        ["Contrato", "Tipo", "Instrumento", "Monto", "Vencimiento", "Estado", "Clasif."],
        [26, 32, 30, 30, 26, 24, 26],
        [[g["contrato"], g["tipo"], g["instrumento"], f"{g['moneda']} {g['monto']}",
          g["fecha_vencimiento"], g["estado"], g["clasificacion"]]
         for g in datos["garantias"]],
    )

    salida = pdf.output()
    return bytes(salida)


def generar_reporte_mensual(
    session: Session, anio: int, mes: int, formato: str = "xlsx"
) -> tuple[bytes, str, str]:
    if formato not in MEDIA:
        raise ValueError("formato debe ser 'xlsx' o 'pdf'")
    datos = datos_reporte(session, anio, mes)
    contenido = reporte_xlsx(datos) if formato == "xlsx" else reporte_pdf(datos)
    nombre = f"reporte_contratos_{datos['periodo']}.{formato}"
    return contenido, nombre, MEDIA[formato]
