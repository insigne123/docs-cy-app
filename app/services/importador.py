"""Importador de contratos desde planilla Excel/CSV (ver docs/03-plantilla-carga-contratos.md).

Carga masiva para migrar contratos existentes y para incorporaciones futuras. No recorre
el flujo: crea cada expediente directamente en su `estado_actual` y deja un único evento
de trazabilidad "carga inicial".
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.enums import (
    EstadoContrato,
    GarantiaEstado,
    GarantiaInstrumento,
    GarantiaTipo,
    HitoEstado,
    HitoTipo,
    LineaContrato,
    Moneda,
    Rol,
    TipoRenovacion,
)
from app.models.contrato import Contrato, Garantia, Hito
from app.models.core import Contraparte, FormatoEstandar, Unidad, Usuario
from app.models.evento import EventoEstado
from app.services.contratos import calcular_requiere_gerencia

_LINEA_MAP = {
    "A": LineaContrato.A_regular,
    "B": LineaContrato.B_autogestionado,
    "C": LineaContrato.C_licitacion,
}
_VERDADEROS = {"si", "sí", "true", "1", "x", "verdadero"}
_CARPETA_DOCUMENTOS = Path("Contratos/Documentos")


class FilaInvalida(Exception):
    pass


@dataclass
class ResultadoImportacion:
    creados: list[str] = field(default_factory=list)
    actualizados: list[str] = field(default_factory=list)
    errores: list[dict] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)

    def resumen(self) -> dict:
        return {
            "creados": self.creados,
            "actualizados": self.actualizados,
            "errores": self.errores,
            "advertencias": self.advertencias,
        }


# --------------------------------------------------------------------- helpers de celda
def _txt(fila: dict, col: str) -> Optional[str]:
    v = fila.get(col)
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _req(fila: dict, col: str) -> str:
    v = _txt(fila, col)
    if v is None:
        raise FilaInvalida(f"falta el campo obligatorio '{col}'")
    return v


def _fecha(valor: Any) -> Optional[date]:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip().replace("/", "-")
    try:
        return date.fromisoformat(s)
    except ValueError as exc:
        raise FilaInvalida(f"fecha inválida '{valor}' (se espera AAAA-MM-DD)") from exc


def _num(valor: Any) -> Optional[Decimal]:
    if valor in (None, ""):
        return None
    try:
        return Decimal(str(valor).strip().replace(" ", ""))
    except (InvalidOperation, ValueError) as exc:
        raise FilaInvalida(f"número inválido '{valor}'") from exc


def _sino(valor: Any, default: bool = False) -> bool:
    if valor in (None, ""):
        return default
    return str(valor).strip().lower() in _VERDADEROS


def _enum(cls, valor: str, col: str):
    try:
        return cls(valor)
    except ValueError as exc:
        opciones = ", ".join(m.value for m in cls)
        raise FilaInvalida(f"'{col}' inválido: '{valor}'. Valores: {opciones}") from exc


# --------------------------------------------------------------------- catálogos
def _unidad(session: Session, nombre: str, res: ResultadoImportacion) -> Unidad:
    from app.enums import UnidadTipo

    u = session.scalars(select(Unidad).where(Unidad.nombre == nombre)).first()
    if u is None:
        u = Unidad(nombre=nombre, tipo=UnidadTipo.solicitante)
        session.add(u)
        session.flush()
        res.advertencias.append(f"Unidad creada: '{nombre}'")
    return u


def _usuario(session: Session, email: str, rol: Rol, res: ResultadoImportacion) -> Usuario:
    u = session.scalars(select(Usuario).where(Usuario.email == email)).first()
    if u is None:
        u = Usuario(nombre=email.split("@")[0], email=email, rol=rol, activo=True)
        session.add(u)
        session.flush()
        res.advertencias.append(f"Usuario creado: '{email}' con rol '{rol.value}' (revisar)")
    return u


def _contraparte(session: Session, razon_social: str, rut: Optional[str], res: ResultadoImportacion) -> Contraparte:
    from app.enums import ContraparteTipo

    c = session.scalars(select(Contraparte).where(Contraparte.razon_social == razon_social)).first()
    if c is None:
        c = Contraparte(razon_social=razon_social, rut=rut, tipo=ContraparteTipo.proveedor)
        session.add(c)
        session.flush()
        res.advertencias.append(f"Contraparte creada: '{razon_social}'")
    return c


def _formato_vigente(session: Session, nombre: str) -> Optional[FormatoEstandar]:
    return session.scalars(
        select(FormatoEstandar)
        .where(FormatoEstandar.nombre == nombre, FormatoEstandar.vigente.is_(True))
        .order_by(FormatoEstandar.version.desc())
    ).first()


def _generar_codigo(session: Session, usados: set[str]) -> str:
    anio = date.today().year
    n = session.query(Contrato).count() + len(usados) + 1
    while True:
        codigo = f"CT-{anio}-{n:04d}"
        if codigo not in usados and session.scalars(
            select(Contrato.id).where(Contrato.codigo == codigo)
        ).first() is None:
            return codigo
        n += 1


# --------------------------------------------------------------------- lectura de archivo
def _leer_hojas(ruta: Path) -> dict[str, list[dict]]:
    if ruta.suffix.lower() == ".csv":
        with ruta.open(encoding="utf-8-sig", newline="") as fh:
            filas = [dict(r) for r in csv.DictReader(fh)]
        return {"contratos": filas}

    from openpyxl import load_workbook

    wb = load_workbook(ruta, read_only=True, data_only=True)
    hojas: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        filas: list[dict] = []
        it = ws.iter_rows(values_only=True)
        try:
            encabezado = [str(c).strip() if c is not None else "" for c in next(it)]
        except StopIteration:
            continue
        for valores in it:
            if valores is None or all(v is None or str(v).strip() == "" for v in valores):
                continue
            filas.append({encabezado[i]: valores[i] for i in range(min(len(encabezado), len(valores)))})
        hojas[ws.title.strip().lower()] = filas
    wb.close()
    return hojas


# --------------------------------------------------------------------- procesamiento
_ORDEN_FECHAS = [
    "fecha_ingreso",
    "fecha_aprobacion_jefatura",
    "fecha_admisibilidad",
    "fecha_visacion",
    "fecha_firma",
    "fecha_integracion",
    "fecha_inicio_vigencia",
    "fecha_fin_vigencia",
]


def _fechas_de(fila: dict, fallback_ingreso: Optional[date]) -> dict[str, Optional[date]]:
    d = {col: _fecha(fila.get(col)) for col in _ORDEN_FECHAS}
    d["fecha_ingreso"] = d["fecha_ingreso"] or fallback_ingreso or date.today()
    presentes = [(n, d[n]) for n in _ORDEN_FECHAS if d[n] is not None]
    for (n1, f1), (n2, f2) in zip(presentes, presentes[1:]):
        if f2 < f1:
            raise FilaInvalida(f"incoherencia de fechas: {n2} ({f2}) es anterior a {n1} ({f1})")
    return d


def _procesar_contrato(
    session: Session, fila: dict, res: ResultadoImportacion, cache: dict, ruta: Path, usados: set[str]
) -> None:
    linea_txt = _req(fila, "linea").upper()
    if linea_txt not in _LINEA_MAP:
        raise FilaInvalida(f"'linea' inválida: '{linea_txt}' (A, B o C)")
    linea = _LINEA_MAP[linea_txt]
    objeto = _req(fila, "objeto")
    estado = _enum(EstadoContrato, _req(fila, "estado_actual"), "estado_actual")

    unidad = _unidad(session, _req(fila, "unidad_solicitante"), res)
    solicitante = _usuario(session, _req(fila, "solicitante_email"), Rol.unidad_solicitante, res)
    contraparte = _contraparte(session, _req(fila, "contraparte_razon_social"), _txt(fila, "contraparte_rut"), res)

    abogado = None
    if _txt(fila, "abogado_email"):
        abogado = _usuario(session, _txt(fila, "abogado_email"), Rol.legal, res)
    administrador = None
    if _txt(fila, "administrador_email"):
        administrador = _usuario(session, _txt(fila, "administrador_email"), Rol.admin_contratos, res)

    formato = None
    if linea == LineaContrato.B_autogestionado:
        nombre_fmt = _req(fila, "formato_estandar")
        formato = _formato_vigente(session, nombre_fmt)
        if formato is None:
            raise FilaInvalida(f"no existe un formato estándar vigente llamado '{nombre_fmt}'")

    moneda = _enum(Moneda, _txt(fila, "moneda"), "moneda") if _txt(fila, "moneda") else None
    monto = _num(fila.get("monto"))
    if monto is not None and moneda is None:
        raise FilaInvalida("hay 'monto' pero falta 'moneda'")
    categoria_txt = _txt(fila, "categoria")
    from app.enums import CategoriaContrato

    categoria = _enum(CategoriaContrato, categoria_txt, "categoria") if categoria_txt else None
    tipo_renov_txt = _txt(fila, "tipo_renovacion")
    tipo_renov = (
        _enum(TipoRenovacion, tipo_renov_txt, "tipo_renovacion")
        if tipo_renov_txt
        else TipoRenovacion.sin_renovacion
    )
    vig_indef = _sino(fila.get("vigencia_indefinida"))
    aviso = int(_num(fila.get("aviso_previo_dias")) or 60)
    monto_ref = _num(fila.get("monto_referencia_clp"))
    req_gar = _sino(fila.get("requiere_garantia"))
    causales = _txt(fila, "causales_termino")

    codigo_externo = _txt(fila, "codigo_externo")
    existente: Optional[Contrato] = None
    if codigo_externo:
        existente = session.scalars(
            select(Contrato).where(Contrato.codigo_externo == codigo_externo)
        ).first()

    # Todo lo que consulta la BD se resuelve ANTES de tocar/crear el Contrato
    # (evita autoflush prematuro de un contrato a medio construir).
    fechas = _fechas_de(fila, existente.fecha_ingreso if existente else None)
    req_ger = calcular_requiere_gerencia(session, monto, moneda, monto_ref)

    campos: dict[str, Any] = dict(
        linea=linea,
        objeto=objeto,
        categoria=categoria,
        estado=estado,
        unidad_solicitante_id=unidad.id,
        solicitante_id=solicitante.id,
        contraparte_id=contraparte.id,
        abogado_id=abogado.id if abogado else None,
        administrador_id=administrador.id if administrador else None,
        formato_id=formato.id if formato else None,
        monto=monto,
        moneda=moneda,
        monto_referencia_clp=monto_ref,
        requiere_garantia=req_gar,
        requiere_aprobacion_gerencia=req_ger,
        vigencia_indefinida=vig_indef,
        tipo_renovacion=tipo_renov,
        aviso_previo_dias=aviso,
        causales_termino=causales,
        **fechas,
    )
    if estado in (EstadoContrato.terminado, EstadoContrato.descartado):
        campos["fecha_cierre"] = fechas["fecha_fin_vigencia"] or date.today()

    nuevo = existente is None
    if nuevo:
        codigo = codigo_externo or _generar_codigo(session, usados)
        usados.add(codigo)
        contrato = Contrato(
            codigo=codigo,
            codigo_externo=codigo_externo,
            estado_desde=datetime.utcnow(),
            **campos,
        )
        session.add(contrato)
    else:
        contrato = existente
        for k, v in campos.items():
            setattr(contrato, k, v)
    session.flush()

    lic_codigo = _txt(fila, "licitacion_codigo")
    if lic_codigo:
        from app.models.licitacion import Licitacion

        lic = session.scalars(select(Licitacion).where(Licitacion.codigo == lic_codigo)).first()
        contrato.licitacion_id = lic.id if lic else None
        if lic is None:
            res.advertencias.append(f"{contrato.codigo}: licitación '{lic_codigo}' no encontrada")

    nombre_doc = _txt(fila, "archivo_contrato")
    if nombre_doc:
        from app.enums import DocumentoTipo
        from app.models.contrato import Documento

        ruta_doc = _CARPETA_DOCUMENTOS / nombre_doc
        if ruta_doc.exists():
            session.add(
                Documento(
                    entidad_tipo="contrato",
                    entidad_id=contrato.id,
                    tipo=DocumentoTipo.contrato_firmado,
                    nombre_archivo=nombre_doc,
                    ruta=str(ruta_doc),
                    cargado_por_id=solicitante.id,
                )
            )
        else:
            res.advertencias.append(f"{contrato.codigo}: no se encontró el archivo '{ruta_doc}'")

    if nuevo:
        session.add(
            EventoEstado(
                entidad_tipo="contrato",
                entidad_id=contrato.id,
                estado_origen=None,
                estado_destino=estado.value,
                fecha=datetime.utcnow(),
                usuario_id=solicitante.id,
                rol_actor=Rol.unidad_solicitante.value,
                comentario=f"Carga inicial desde {ruta.name}",
            )
        )
        res.creados.append(contrato.codigo)
    else:
        res.actualizados.append(contrato.codigo)

    cache[contrato.codigo_externo or contrato.codigo] = contrato


def _buscar_contrato(session: Session, clave: str, cache: dict) -> Optional[Contrato]:
    if clave in cache:
        return cache[clave]
    return session.scalars(select(Contrato).where(Contrato.codigo_externo == clave)).first() or (
        session.scalars(select(Contrato).where(Contrato.codigo == clave)).first()
    )


def _procesar_garantia(session: Session, fila: dict, res: ResultadoImportacion, cache: dict) -> None:
    clave = _req(fila, "contrato_codigo_externo")
    contrato = _buscar_contrato(session, clave, cache)
    if contrato is None:
        raise FilaInvalida(f"no existe el contrato '{clave}'")
    session.add(
        Garantia(
            entidad_tipo="contrato",
            entidad_id=contrato.id,
            tipo=_enum(GarantiaTipo, _req(fila, "tipo"), "tipo"),
            instrumento=_enum(GarantiaInstrumento, _req(fila, "instrumento"), "instrumento"),
            emisor=_txt(fila, "emisor"),
            numero=_txt(fila, "numero"),
            monto=_num(_req(fila, "monto")),
            moneda=_enum(Moneda, _req(fila, "moneda"), "moneda"),
            fecha_emision=_fecha(_req(fila, "fecha_emision")),
            fecha_vencimiento=_fecha(_req(fila, "fecha_vencimiento")),
            estado=_enum(GarantiaEstado, _txt(fila, "estado") or "vigente", "estado"),
            glosa=_txt(fila, "glosa"),
        )
    )


def _procesar_hito(session: Session, fila: dict, res: ResultadoImportacion, cache: dict) -> None:
    clave = _req(fila, "contrato_codigo_externo")
    contrato = _buscar_contrato(session, clave, cache)
    if contrato is None:
        raise FilaInvalida(f"no existe el contrato '{clave}'")
    responsable = None
    if _txt(fila, "responsable_email"):
        responsable = _usuario(session, _txt(fila, "responsable_email"), Rol.tecnica, res)
    session.add(
        Hito(
            contrato_id=contrato.id,
            tipo=_enum(HitoTipo, _req(fila, "tipo"), "tipo"),
            nombre=_req(fila, "nombre"),
            fecha_planificada=_fecha(_req(fila, "fecha_planificada")),
            fecha_real=_fecha(fila.get("fecha_real")),
            estado=_enum(HitoEstado, _txt(fila, "estado") or "pendiente", "estado"),
            responsable_id=responsable.id if responsable else None,
            notas=_txt(fila, "notas"),
        )
    )


def _procesar_hoja(
    filas: Iterable[dict], hoja: str, procesar, session: Session, res: ResultadoImportacion, *args
) -> None:
    for i, fila in enumerate(filas, start=2):  # fila 1 = encabezado
        sp = session.begin_nested()
        try:
            procesar(session, fila, res, *args)
            sp.commit()
        except FilaInvalida as exc:
            sp.rollback()
            res.errores.append({"hoja": hoja, "fila": i, "motivo": str(exc)})
        except Exception as exc:  # noqa: BLE001 - se reporta y se continúa
            sp.rollback()
            res.errores.append({"hoja": hoja, "fila": i, "motivo": f"error inesperado: {exc}"})


def importar_planilla_contratos(session: Session, ruta: Path | str) -> ResultadoImportacion:
    ruta = Path(ruta)
    res = ResultadoImportacion()
    hojas = _leer_hojas(ruta)
    if "contratos" not in hojas:
        res.errores.append({"hoja": "-", "fila": 0, "motivo": "no se encontró la hoja 'Contratos'"})
        return res

    cache: dict = {}
    usados: set[str] = set()
    _procesar_hoja(hojas["contratos"], "Contratos", _procesar_contrato, session, res, cache, ruta, usados)
    if "garantias" in hojas:
        _procesar_hoja(hojas["garantias"], "Garantias", _procesar_garantia, session, res, cache)
    if "hitos" in hojas:
        _procesar_hoja(hojas["hitos"], "Hitos", _procesar_hito, session, res, cache)
    return res
