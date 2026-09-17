"""Reporte mensual descargable (XLSX y PDF)."""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.services.reportes import datos_reporte, generar_reporte_mensual

# Mes en que vence C-VENCIDO (hoy - 10 días); su corte de fin de mes lo deja "vencido".
_MES_VENCIDO = date.today() - timedelta(days=10)


def test_datos_reporte(session, cartera):
    d = datos_reporte(session, _MES_VENCIDO.year, _MES_VENCIDO.month)
    assert d["periodo"] == _MES_VENCIDO.strftime("%Y-%m")
    assert len(d["detalle"]) == 5
    assert {"vigentes", "vencidos", "por_vencer"} <= set(d["indicadores"])
    assert any(a["tipo"] == "contrato_vencido" for a in d["alertas"])
    assert "C-VENCIDO" in [c["codigo"] for c in d["vencen_en_el_mes"]]


def test_reporte_xlsx(session, cartera):
    contenido, nombre, media = generar_reporte_mensual(
        session, _MES_VENCIDO.year, _MES_VENCIDO.month, "xlsx"
    )
    assert nombre.endswith(".xlsx")
    assert "spreadsheet" in media
    wb = load_workbook(BytesIO(contenido))
    assert {"Resumen", "Detalle", "Alertas", "Garantias"} <= set(wb.sheetnames)
    assert wb["Detalle"].max_row == 6  # encabezado + 5 contratos


def test_reporte_pdf(session, cartera):
    contenido, nombre, media = generar_reporte_mensual(
        session, _MES_VENCIDO.year, _MES_VENCIDO.month, "pdf"
    )
    assert nombre.endswith(".pdf")
    assert media == "application/pdf"
    assert contenido[:4] == b"%PDF"
    assert len(contenido) > 800


def test_formato_invalido(session, cartera):
    import pytest

    with pytest.raises(ValueError):
        generar_reporte_mensual(session, 2026, 3, "docx")


def test_endpoint_reporte(api, cartera):
    periodo = _MES_VENCIDO.strftime("%Y-%m")
    r = api.get("/panel/reporte", params={"periodo": periodo, "formato": "xlsx"})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert "spreadsheet" in r.headers["content-type"]

    r2 = api.get("/panel/reporte", params={"periodo": periodo, "formato": "pdf"})
    assert r2.status_code == 200
    assert r2.content[:4] == b"%PDF"

    assert api.get("/panel/reporte", params={"periodo": "malo", "formato": "xlsx"}).status_code == 400
