"""Importador de planillas Excel/CSV."""
from __future__ import annotations

from openpyxl import Workbook

from app.models.contrato import Contrato, Garantia, Hito
from app.services.importador import importar_planilla_contratos

HDR_CONTRATOS = [
    "codigo_externo", "linea", "objeto", "categoria", "unidad_solicitante", "solicitante_email",
    "contraparte_razon_social", "contraparte_rut", "contraparte_contacto_email", "abogado_email",
    "administrador_email", "formato_estandar", "licitacion_codigo", "estado_actual", "monto",
    "moneda", "requiere_garantia", "fecha_ingreso", "fecha_aprobacion_jefatura", "fecha_admisibilidad",
    "fecha_visacion", "fecha_firma", "fecha_integracion", "fecha_inicio_vigencia", "fecha_fin_vigencia",
    "vigencia_indefinida", "tipo_renovacion", "aviso_previo_dias", "causales_termino",
    "archivo_contrato", "observaciones",
]
HDR_GARANTIAS = [
    "contrato_codigo_externo", "tipo", "instrumento", "emisor", "numero", "monto", "moneda",
    "fecha_emision", "fecha_vencimiento", "estado", "glosa",
]
HDR_HITOS = [
    "contrato_codigo_externo", "tipo", "nombre", "fecha_planificada", "fecha_real", "estado",
    "responsable_email", "notas",
]


def _fila_contrato(codigo, estado="vigente", linea="A", monto="1000000", **over):
    base = {h: "" for h in HDR_CONTRATOS}
    base.update(
        codigo_externo=codigo, linea=linea, objeto=f"Objeto {codigo}", categoria="servicio",
        unidad_solicitante="Operaciones", solicitante_email="jperez@empresa.cl",
        contraparte_razon_social="Proveedor SpA", estado_actual=estado, monto=monto, moneda="CLP",
        fecha_ingreso="2025-01-10", fecha_fin_vigencia="2026-12-31",
        vigencia_indefinida="no", tipo_renovacion="automatica", aviso_previo_dias="60",
    )
    base.update(over)
    return [base[h] for h in HDR_CONTRATOS]


def _construir_xlsx(ruta, contratos, garantias=(), hitos=()):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Contratos")
    ws.append(HDR_CONTRATOS)
    for f in contratos:
        ws.append(f)
    if garantias:
        wg = wb.create_sheet("Garantias")
        wg.append(HDR_GARANTIAS)
        for f in garantias:
            wg.append(f)
    if hitos:
        wh = wb.create_sheet("Hitos")
        wh.append(HDR_HITOS)
        for f in hitos:
            wh.append(f)
    wb.save(ruta)


def test_importar_xlsx_crea_actualiza_y_reporta_errores(session, tmp_path):
    ruta = tmp_path / "carga.xlsx"
    contratos = [
        _fila_contrato("SC-1"),
        _fila_contrato("SC-2", estado="elaboracion", monto="60000000"),
        _fila_contrato("SC-3", estado="ESTADO_INVALIDO"),
    ]
    garantias = [
        ["SC-1", "fiel_cumplimiento", "boleta_bancaria", "Banco Estado", "007", "500000", "CLP",
         "2025-02-01", "2026-06-30", "vigente", "glosa"],
    ]
    hitos = [
        ["SC-1", "renovacion", "Aviso renovación", "2026-10-01", "", "pendiente", "", ""],
    ]
    _construir_xlsx(ruta, contratos, garantias, hitos)

    res = importar_planilla_contratos(session, ruta)
    session.commit()

    assert sorted(res.creados) == ["SC-1", "SC-2"]
    assert len(res.errores) == 1
    assert res.errores[0]["fila"] == 4  # 3ª fila de datos

    c1 = session.query(Contrato).filter_by(codigo_externo="SC-1").one()
    assert c1.estado.value == "vigente"
    assert c1.requiere_aprobacion_gerencia is False
    c2 = session.query(Contrato).filter_by(codigo_externo="SC-2").one()
    assert c2.estado.value == "elaboracion"
    assert c2.requiere_aprobacion_gerencia is True  # 60.000.000 CLP > umbral (50M)

    assert session.query(Garantia).filter_by(entidad_tipo="contrato", entidad_id=c1.id).count() == 1
    assert session.query(Hito).filter_by(contrato_id=c1.id).count() == 1
    assert session.query(Contrato).count() == 2

    # Re-importar: actualiza, no duplica
    res2 = importar_planilla_contratos(session, ruta)
    session.commit()
    assert sorted(res2.actualizados) == ["SC-1", "SC-2"]
    assert res2.creados == []
    assert session.query(Contrato).count() == 2


def test_importar_csv(session, tmp_path):
    ruta = tmp_path / "carga.csv"
    lineas = [",".join(HDR_CONTRATOS)]
    fila = _fila_contrato("CSV-1")
    lineas.append(",".join(str(x) for x in fila))
    ruta.write_text("\n".join(lineas), encoding="utf-8")

    res = importar_planilla_contratos(session, ruta)
    session.commit()
    assert res.creados == ["CSV-1"]
    assert session.query(Contrato).filter_by(codigo_externo="CSV-1").one().estado.value == "vigente"


def test_endpoint_importacion(api, session, tmp_path):
    ruta = tmp_path / "endpoint.xlsx"
    _construir_xlsx(ruta, [_fila_contrato("EP-1"), _fila_contrato("EP-2")])
    r = api.post("/importaciones/contratos", json={"archivo": "endpoint.xlsx", "carpeta": str(tmp_path)})
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert sorted(cuerpo["creados"]) == ["EP-1", "EP-2"]
    assert cuerpo["errores"] == []
    assert len(api.get("/contratos", params={"estado": "vigente"}).json()) == 2


def test_endpoint_importacion_archivo_inexistente(api, tmp_path):
    r = api.post("/importaciones/contratos", json={"archivo": "no_existe.xlsx", "carpeta": str(tmp_path)})
    assert r.status_code == 404
